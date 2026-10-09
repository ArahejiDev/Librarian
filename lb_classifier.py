"""
LB- document classifier (Codiv System One models)

Watches a folder. Whenever a document whose name starts with "LB-" shows up,
it is classified (work / leisure / other) and moved to OUTPUT_DIR/<category>/.
If the model is unsure (confidence < CONFIDENCE_THRESHOLD) the file goes to
OUTPUT_DIR/manual_review/. Files that do NOT start with the prefix are ignored
and left untouched.

Formats: .txt .md .eml .log .csv .json  (+ .pdf with pypdf, + .docx with python-docx)

Dependencies:  pip install requests pypdf python-docx

Environment variables:
  CODIV_API_KEY         (required)
  INPUT_DIR             watched folder                         [./input]
  OUTPUT_DIR            results folder                         [./output]
  PREFIX                filename prefix that triggers the job  [LB-]
  MODEL                 jevk5-0.2 | laya-1.0 | verdict-1.4 | clm-v0.1   [jevk5-0.2]
  CONFIDENCE_THRESHOLD  below this -> manual review            [0.5]
  POLL_INTERVAL         seconds between folder scans           [5]
  MAX_CHARS             characters sent to the model           [8000]
  CATEGORIES_FILE       JSON {"category": "description"}       [optional]

Usage:
  python lb_classifier.py            # keep watching
  python lb_classifier.py --once     # process what is there and exit (cron / tests)
"""
import csv
import json
import logging
import os
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import requests

# Configuration
API_URL = os.getenv("CODIV_API_URL", "https://api.codiv.ai/v1/systemone")
API_KEY = os.getenv("CODIV_API_KEY", "")
MODEL = os.getenv("MODEL", "jevk5-0.2")  # text-only model
THRESHOLD = float(os.getenv("CONFIDENCE_THRESHOLD", "0.5"))
MAX_CHARS = int(os.getenv("MAX_CHARS", "8000"))
POLL_INTERVAL = float(os.getenv("POLL_INTERVAL", "5"))
MIN_AGE = 3  
PREFIX = os.getenv("PREFIX", "LB-")

INPUT_DIR = Path(os.getenv("INPUT_DIR", "./input"))
OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR", "./output"))
REVIEW_DIR = OUTPUT_DIR / "manual_review"
ERRORS_DIR = OUTPUT_DIR / "errors"
REGISTRY = OUTPUT_DIR / "registry.csv"

CATEGORIES = {
    "work": "Professional content: meetings, reports, clients, projects, company invoices, incidents, work tasks",
    "leisure": "Free time: TV shows, movies, games, travel, sports, hobbies, plans with friends or family",
    "other": "Anything that is clearly neither work nor leisure: personal errands, shopping, bills, loose notes",
}
if os.getenv("CATEGORIES_FILE"):
    CATEGORIES = json.loads(Path(os.environ["CATEGORIES_FILE"]).read_text(encoding="utf-8"))

QUESTIONS = {
    "category": {
        "type": "choice",
        "instructions": "What kind of content is this?",
        "criteria": CATEGORIES,
    }
}

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("lb-classifier")

TEXT_EXTENSIONS = {".txt", ".md", ".eml", ".log", ".csv", ".json"}
EXTENSIONS = set(TEXT_EXTENSIONS)
try:
    import pypdf  # noqa: F401
    EXTENSIONS.add(".pdf")
except ImportError:
    log.warning("pypdf not installed: .pdf files will be ignored (pip install pypdf)")
try:
    import docx  # noqa: F401
    EXTENSIONS.add(".docx")
except ImportError:
    log.warning("python-docx not installed: .docx files will be ignored (pip install python-docx)")


class ConfigError(Exception):
    """Invalid API key or no credit left: retrying makes no sense."""


# Reading documents 
def extract_text(path):
    ext = path.suffix.lower()
    if ext in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        reader = pypdf.PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:20])
    if ext == ".docx":
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported format: {ext}")


# API 
def retry_delay(resp, attempt):
    try:
        return int(resp.headers.get("Retry-After", 2 * attempt))
    except ValueError:
        return 2 * attempt


def classify(text):
    """Return. Retries when the API is busy."""
    for attempt in range(1, 4):
        resp = requests.post(
            API_URL,
            headers={
                "Authorization": f"Bearer {API_KEY}",
                "Content-Type": "application/json",
            },
            json={"model": MODEL, "state": text[:MAX_CHARS], "questions": QUESTIONS},
            timeout=60,
        )
        if resp.status_code in (401, 402):
            raise ConfigError(f"HTTP {resp.status_code}: check your API key / credit")
        if resp.status_code in (429, 502, 503, 504, 529):
            delay = retry_delay(resp, attempt)
            log.warning("API busy (HTTP %s), retrying in %ss", resp.status_code, delay)
            time.sleep(delay)
            continue
        resp.raise_for_status()
        answer = resp.json()["answers"]["category"]
        return answer["choice"], answer["confidence"], answer.get("probabilities", {})
    raise RuntimeError("API unavailable after 3 attempts")


# Files 
def move_file(path, folder):
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / path.name
    if target.exists():
        target = folder / f"{path.stem}_{int(time.time())}{path.suffix}"
    shutil.move(str(path), str(target))
    return target


def write_registry(filename, category, confidence, target):
    is_new = not REGISTRY.exists()
    REGISTRY.parent.mkdir(parents=True, exist_ok=True)
    with open(REGISTRY, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if is_new:
            writer.writerow(["date", "file", "category", "confidence", "destination"])
        writer.writerow([datetime.now().isoformat(timespec="seconds"), filename,
                         category, f"{confidence:.3f}", target])


def process(path):
    text = extract_text(path)

    if not text.strip() or "\x00" in text:
        target = move_file(path, REVIEW_DIR)
        write_registry(path.name, "no_text", 0, target)
        log.info("[REVIEW] %s -> manual_review (no readable text)", path.name)
        return

    category, confidence, probs = classify(text)
    spread = " ".join(f"{k}={v:.0%}" for k, v in sorted(probs.items(), key=lambda kv: -kv[1]))

    if confidence < THRESHOLD:
        target = move_file(path, REVIEW_DIR)
        write_registry(path.name, category, confidence, target)
        log.info("[REVIEW] %s -> manual_review (%s, conf %.2f | %s)", path.name, category, confidence, spread)
    else:
        target = move_file(path, OUTPUT_DIR / category)
        write_registry(path.name, category, confidence, target)
        log.info("[OK    ] %s -> %s (conf %.2f | %s)", path.name, category, confidence, spread)


def is_candidate(path, ignored):
    """True if it is a file with the right prefix and format, and it is no longer being written."""
    if not path.is_file():
        return False
    if not path.name.upper().startswith(PREFIX.upper()):
        if path.name not in ignored:
            ignored.add(path.name)
            log.info("[SKIP  ] %s (does not start with %s)", path.name, PREFIX)
        return False
    if path.suffix.lower() not in EXTENSIONS:
        if path.name not in ignored:
            ignored.add(path.name)
            log.info("[SKIP  ] %s (unsupported format)", path.name)
        return False
    return time.time() - path.stat().st_mtime >= MIN_AGE


def scan_once(ignored):
    """Process whatever is in INPUT_DIR."""
    for path in sorted(INPUT_DIR.iterdir()):
        if not is_candidate(path, ignored):
            continue
        try:
            process(path)
        except ConfigError as e:
            log.error("%s. Pausing for 60s.", e)
            time.sleep(60)
            return
        except requests.HTTPError as e:
            move_file(path, ERRORS_DIR)
            log.error("[ERROR ] %s -> errors (%s)", path.name, e)
        except (requests.RequestException, RuntimeError) as e:
            log.warning("[RETRY ] %s stays in the input folder: %s", path.name, e)
        except Exception as e:
            move_file(path, ERRORS_DIR)
            log.exception("[ERROR ] %s -> errors (%s)", path.name, e)


def main():
    if not API_KEY:
        sys.exit("Missing environment variable CODIV_API_KEY")
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    log.info("Watching %s | prefix=%s | model=%s | threshold=%.2f | formats=%s",
             INPUT_DIR, PREFIX, MODEL, THRESHOLD, ",".join(sorted(EXTENSIONS)))

    ignored = set()
    if "--once" in sys.argv:
        time.sleep(MIN_AGE)  # let freshly copied files settle
        scan_once(ignored)
        return

    while True:
        scan_once(ignored)
        time.sleep(POLL_INTERVAL)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        log.info("Stopped.")
